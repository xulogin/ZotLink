# -*- coding: utf-8 -*-
"""Build a Zotero-linked Word document from a Markdown content file.

Reads (relative to the project root):
  content/<name>.md          first .md found, or the path given on the CLI
  data/references.json       CSL JSON produced by pandoc
  reports/zot_match_index.json

Writes:
  output/<name>.docx

Markdown syntax supported:
  ---
  title: <document title>
  ---

  # Heading 1
  ## Heading 2

  Plain paragraph text, may contain [[cite:key1;key2]] markers. Consecutive
  lines join into one paragraph (no separator between two CJK characters,
  a space otherwise), blank line ends it.

  (1) **Bold lead-in**: rest of the item [[cite:key]] ...

  %% a line starting with %% is a comment and is never rendered

  [[bibliography: 参考文献 / References]]
      inserts the Zotero CSL_BIBLIOGRAPHY placeholder, optionally preceded
      by a heading with the text after the colon.

Typography (Times New Roman + 宋体/黑体, 1.5 line spacing, 0.74 cm first-line
indent) is defined in the add_* functions — edit there for a different layout.

Run:
  python scripts/build_doc.py [--project DIR] [content/my.md]
"""
from __future__ import annotations

import json
import random
import re
import string
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
import zot_env  # noqa: E402

try:
    from docx import Document
    from docx.enum.text import WD_LINE_SPACING, WD_PARAGRAPH_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor
except ImportError:
    sys.exit("FATAL: python-docx is not installed — pip install python-docx")

PROJ, ARGS = zot_env.open_project()
PROJ.ensure_dirs()

# ---------------------------------------------------------------------------
# Resolve content file
# ---------------------------------------------------------------------------
if ARGS:
    md_path = Path(ARGS[0])
    if not md_path.is_absolute():
        md_path = (PROJ.root / md_path) if (PROJ.root / md_path).exists() else md_path
    md_path = md_path.resolve()
else:
    candidates = sorted(PROJ.content.glob("*.md"))
    if not candidates:
        sys.exit("FATAL: no .md files in {}".format(PROJ.content))
    md_path = candidates[0]
if not md_path.exists():
    sys.exit("FATAL: {} not found".format(md_path))

OUT_PATH = PROJ.output / (md_path.stem + ".docx")

# ---------------------------------------------------------------------------
# Load CSL registry + Zotero match index
# ---------------------------------------------------------------------------
for required in (PROJ.csl_json, PROJ.match_index):
    if not required.exists():
        sys.exit("FATAL: {} not found — run scripts/run.py first".format(required))

with open(PROJ.csl_json, "r", encoding="utf-8") as f:
    CSL_REGISTRY = dict((item["id"], item) for item in json.load(f))
with open(PROJ.match_index, "r", encoding="utf-8") as f:
    MATCH_INDEX = json.load(f)

print("[content]  {}".format(md_path))
print("[csl]      {} entries".format(len(CSL_REGISTRY)))
print("[match]    {} URI mappings".format(len(MATCH_INDEX)))


def _check_keys(keys):
    unknown = [k for k in keys if k not in CSL_REGISTRY]
    unmatched = [k for k in keys if k in CSL_REGISTRY and k not in MATCH_INDEX]
    if unknown:
        sys.exit("FATAL: citation key(s) not in references.bib: {}\n"
                 "       add the BibTeX entry, then re-run scripts/run.py"
                 .format(", ".join(unknown)))
    if unmatched:
        sys.exit("FATAL: citation key(s) not found in your Zotero library: {}\n"
                 "       import reports/references_missing.bib into Zotero, "
                 "then re-run scripts/run.py".format(", ".join(unmatched)))


# ---------------------------------------------------------------------------
# Citation display text (Harvard; Zotero rewrites it on Refresh anyway)
# ---------------------------------------------------------------------------
def _year_of(csl):
    try:
        return str(csl["issued"]["date-parts"][0][0])
    except (KeyError, IndexError):
        return "n.d."


def _author_string(csl):
    authors = csl.get("author") or csl.get("editor") or []
    if not authors:
        return csl.get("title", "Anon.")
    families = [a.get("family", a.get("literal", "?")) for a in authors]
    if len(families) == 1:
        return families[0]
    if len(families) == 2:
        return "{} and {}".format(families[0], families[1])
    return "{} et al.".format(families[0])


def _format_inline(keys):
    by_author = {}
    order = []
    for k in keys:
        csl = CSL_REGISTRY[k]
        auth = _author_string(csl)
        if auth not in by_author:
            by_author[auth] = []
            order.append(auth)
        by_author[auth].append(_year_of(csl))
    parts = []
    for auth in order:
        years = sorted(set(by_author[auth]))
        parts.append("{}, {}".format(auth, ", ".join(years)))
    return "(" + "; ".join(parts) + ")"


# ---------------------------------------------------------------------------
# Zotero field helpers
# ---------------------------------------------------------------------------
_CITATION_COUNTER = 0


def _rand_id(n=10):
    return "".join(random.choices(string.ascii_letters + string.digits, k=n))


def _zotero_citation_json(keys, display_text):
    """CSL JSON carrying real `uris`, so Zotero matches by URI on Refresh."""
    global _CITATION_COUNTER
    _CITATION_COUNTER += 1
    citation_items = []
    for k in keys:
        csl = CSL_REGISTRY[k]
        match = MATCH_INDEX[k]
        citation_items.append({
            "id": match["itemID"],
            "uris": [match["uri"]],
            "itemData": csl,
        })
    payload = {
        "citationID": _rand_id(),
        "properties": {
            "formattedCitation": display_text,
            "plainCitation": display_text,
            "noteIndex": 0,
        },
        "citationItems": citation_items,
        "schema": "https://github.com/citation-style-language/schema/raw/master/csl-citation.json",
    }
    return json.dumps(payload, ensure_ascii=False, separators=(",", ":"))


def _make_rpr_eastasia():
    rpr = OxmlElement("w:rPr")
    rfonts = OxmlElement("w:rFonts")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")
    rfonts.set(qn("w:eastAsia"), "宋体")
    rpr.append(rfonts)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "22")
    rpr.append(sz)
    return rpr


def _append_text_run(paragraph, text):
    if not text:
        return
    r = OxmlElement("w:r")
    r.append(_make_rpr_eastasia())
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r.append(t)
    paragraph._p.append(r)


def _append_field(paragraph, instr, display):
    """Emit a Word complex field: begin / instrText / separate / text / end."""
    for fld_type, body in [
        ("begin", None), ("instr", instr), ("separate", None),
        ("text", display), ("end", None),
    ]:
        r = OxmlElement("w:r")
        r.append(_make_rpr_eastasia())
        if fld_type == "instr":
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = body
            r.append(it)
        elif fld_type == "text":
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = body
            r.append(t)
        else:
            fc = OxmlElement("w:fldChar")
            fc.set(qn("w:fldCharType"), fld_type)
            r.append(fc)
        paragraph._p.append(r)


def _append_zotero_citation(paragraph, keys):
    display = _format_inline(keys)
    instr = " ADDIN ZOTERO_ITEM CSL_CITATION {} RND {} ".format(
        _zotero_citation_json(keys, display), _rand_id(10))
    _append_field(paragraph, instr, display)


def _append_zotero_bibliography(paragraph, placeholder):
    bibl_payload = {"uncited": [], "omitted": [], "custom": []}
    instr = " ADDIN ZOTERO_BIBL {} CSL_BIBLIOGRAPHY ".format(
        json.dumps(bibl_payload, ensure_ascii=False))
    _append_field(paragraph, instr, placeholder)


# ---------------------------------------------------------------------------
# Document element builders (typography)
# ---------------------------------------------------------------------------
_CITE_RE = re.compile(r"\[\[cite:([^\]]+)\]\]")

BIBL_PLACEHOLDER = "【在 Word 中点 Zotero → Refresh，此处将生成参考文献列表】"


def add_title(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(text)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r.font.size = Pt(16)
    r.bold = True


def add_heading_zh(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Times New Roman"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
        run.font.color.rgb = RGBColor(0, 0, 0)


def add_marked_paragraph(doc, text, indent=True):
    p = doc.add_paragraph()
    p.alignment = WD_PARAGRAPH_ALIGNMENT.JUSTIFY
    pf = p.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    if indent:
        pf.first_line_indent = Cm(0.74)
    pf.space_after = Pt(6)

    cursor = 0
    for m in _CITE_RE.finditer(text):
        _append_text_run(p, text[cursor:m.start()])
        keys = [k.strip() for k in m.group(1).split(";") if k.strip()]
        _check_keys(keys)
        _append_zotero_citation(p, keys)
        cursor = m.end()
    _append_text_run(p, text[cursor:])


def add_bullet(doc, prefix, bold_segment, rest):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.first_line_indent = Cm(0)
    pf.left_indent = Cm(0.74)
    pf.space_after = Pt(4)

    r1 = p.add_run(prefix)
    r1.font.name = "Times New Roman"
    r1._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    r1.bold = True

    r2 = p.add_run(bold_segment)
    r2.font.name = "Times New Roman"
    r2._element.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    r2.bold = True

    cursor = 0
    for m in _CITE_RE.finditer(rest):
        if rest[cursor:m.start()]:
            r = p.add_run(rest[cursor:m.start()])
            r.font.name = "Times New Roman"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        keys = [k.strip() for k in m.group(1).split(";") if k.strip()]
        _check_keys(keys)
        _append_zotero_citation(p, keys)
        cursor = m.end()
    if rest[cursor:]:
        r = p.add_run(rest[cursor:])
        r.font.name = "Times New Roman"
        r._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


# ---------------------------------------------------------------------------
# Markdown parser
# ---------------------------------------------------------------------------
_HEADING_RE = re.compile(r"^(#{1,6})\s+(.+?)\s*$")
_BULLET_RE = re.compile(r"^\((\d+)\)\s+\*\*([^*]+?)\*\*\s*(.+)$")
_BIBL_RE = re.compile(r"^\[\[bibliography(?::\s*(.+?))?\]\]\s*$")


def _is_cjk(ch):
    o = ord(ch)
    return (0x2E80 <= o <= 0x9FFF        # radicals, kana, CJK ideographs
            or 0x3000 <= o <= 0x303F     # CJK punctuation
            or 0xFF00 <= o <= 0xFFEF)    # fullwidth forms


def _join_lines(lines):
    """Join soft-wrapped lines: nothing between two CJK chars, else a space."""
    out = ""
    for line in lines:
        if not out:
            out = line
            continue
        if _is_cjk(out[-1]) and _is_cjk(line[0]):
            out += line
        else:
            out += " " + line
    return out


def parse_markdown(text):
    """Return (frontmatter dict, list of element dicts)."""
    lines = text.split("\n")
    meta = {}

    if lines and lines[0].strip() == "---":
        i = 1
        while i < len(lines) and lines[i].strip() != "---":
            line = lines[i]
            if ":" in line:
                k, v = line.split(":", 1)
                meta[k.strip()] = v.strip()
            i += 1
        lines = lines[i + 1:]

    elements = []
    buf = []

    def flush_paragraph():
        if buf:
            elements.append({"type": "paragraph", "text": _join_lines(buf)})
            del buf[:]

    for line in lines:
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            continue

        # %% comment line — never rendered, safe place for notes and for
        # showing citation syntax without it being parsed as a citation
        if stripped.startswith("%%"):
            continue

        m = _HEADING_RE.match(stripped)
        if m:
            flush_paragraph()
            elements.append({"type": "heading", "level": len(m.group(1)), "text": m.group(2)})
            continue

        m = _BIBL_RE.match(stripped)
        if m:
            flush_paragraph()
            elements.append({"type": "bibliography", "heading": m.group(1)})
            continue

        m = _BULLET_RE.match(stripped)
        if m:
            flush_paragraph()
            elements.append({
                "type": "bullet",
                "prefix": "({}) ".format(m.group(1)),
                "bold": m.group(2).strip(),
                "rest": m.group(3),
            })
            continue

        buf.append(stripped)

    flush_paragraph()
    return meta, elements


# ---------------------------------------------------------------------------
# Render
# ---------------------------------------------------------------------------
meta, elements = parse_markdown(md_path.read_text(encoding="utf-8"))

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

style = doc.styles["Normal"]
style.font.name = "Times New Roman"
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
style.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
style.paragraph_format.first_line_indent = Cm(0.74)
style.paragraph_format.space_after = Pt(6)

if "title" in meta:
    add_title(doc, meta["title"])

bibl_count = 0
for el in elements:
    t = el["type"]
    if t == "heading":
        add_heading_zh(doc, el["text"], level=el["level"])
    elif t == "paragraph":
        add_marked_paragraph(doc, el["text"])
    elif t == "bullet":
        add_bullet(doc, el["prefix"], el["bold"], el["rest"])
    elif t == "bibliography":
        if el["heading"]:
            add_heading_zh(doc, el["heading"], level=1)
        bibl_p = doc.add_paragraph()
        bibl_p.paragraph_format.first_line_indent = Cm(0)
        _append_zotero_bibliography(bibl_p, BIBL_PLACEHOLDER)
        bibl_count += 1

doc.save(OUT_PATH)
print("\n[saved]    {}".format(OUT_PATH))
print("[fields]   {} CSL_CITATION + {} CSL_BIBLIOGRAPHY".format(_CITATION_COUNTER, bibl_count))
if bibl_count == 0:
    print("[warn]     no [[bibliography: ...]] marker in the markdown — "
          "Refresh will format citations but produce no reference list")
